from copy import deepcopy
from datetime import datetime,timezone
import gzip
import hashlib
import io
import json
from pathlib import Path
import re
import unicodedata
import yaml
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
FIXTURES=Path(__file__).parent/'fixtures'
DOC=yaml.safe_load((ROOT/'formal_selection_quote_v18_2_1_http_api.yml').read_text())
NODES={n['id']:n for n in DOC['workflow']['graph']['nodes']}


def load_node(nid):
    namespace={};exec(NODES[nid]['data']['code'],namespace);return namespace


def verified_line(raw,line,chunk):
    start=raw.index(line)
    return {'text':line,'document_hash':hashlib.sha256(raw.encode()).hexdigest(),'chunk_id':chunk,
            'source_verified':True,'source_start':start,'source_end':start+len(line),'source_spans':[[start,start+len(line)]]}


def identity(value):return re.sub(r'\s+','',unicodedata.normalize('NFKC',str(value))).casefold()


def case_requirements():
    """Reconstruct deterministic post-extraction inputs from the uploaded workbook and actual PDF text."""
    context=json.loads((FIXTURES/'v18_1_bom_context.json').read_text())
    technical=(FIXTURES/'technical_case.txt').read_text();cad=(FIXTURES/'cad_case.txt').read_text()
    root=load_node('1800000000007')
    rows=[]
    for item in context['items']:
        row=deepcopy(item)
        row['mapping_status']=item['upstream_mapping_status']
        row['source_evidence']=deepcopy(item['document_evidence'])
        row['handling_mode']='PRODUCT_SELECTION';row['retrieval_required']=True
        row['missing_parameters']=[]
        model=identity(item['shown_model']);role=identity(item['component_role'])
        line=next(line.strip() for line in technical.splitlines() if model in identity(line) and role in identity(line))
        ev=verified_line(technical,line,'CASE-PROTOCOL')
        protocol_requirement={'source_evidence':{'protocol':[ev]}}
        root['restore_peer_protocol_evidence'](row,protocol_requirement)
        rows.append(row)
    cabinet_line=next(line.strip() for line in cad.splitlines() if '双门' in line)
    drawing={'scheme_groups':[{'scheme_group_id':rows[0]['scheme_group_id'],'evidence':[verified_line(cad,cabinet_line,'CASE-CAD')]}]}
    payload={'component_requirements':rows}
    issues=[];root['preserve_source_units_and_structure'](payload,drawing,{},issues)
    return rows,issues


def real_candidates():
    qf=json.loads(gzip.decompress((FIXTURES/'live_qf_search.json.gz').read_bytes()))
    other=json.loads((ROOT/'mongodb_api/tests/fixtures/core_products.json').read_text())
    return qf,other


def selected_case():
    rows,issues=case_requirements()
    assert not issues,issues
    qf,other=real_candidates();selector=load_node('2000000000024')
    outputs=[]
    for req in rows:
        candidates=qf if req['symbol']=='QF' else [p for p in other if req['shown_model'] in p['description']]
        result=selector['main'](json.dumps(req,ensure_ascii=False),json.dumps(req,ensure_ascii=False),candidates)
        outputs.append(result['item_json'])
    aggregator=load_node('2000000000025')
    result=aggregator['main'](outputs,aggregator['wire_pack'](rows))
    return result,rows


def fill_output(out,key='payload_json'):
    workbook=load_workbook(io.BytesIO(__import__('base64').b64decode(out['template_base64'])))
    payload=json.loads(out[key])
    for key,value in payload.get('cells',{}).items():
        sheet,address=key.rsplit('!',1)
        workbook[sheet][address]=value
    output=io.BytesIO();workbook.save(output)
    return output.getvalue()


def dify_extract(blob):
    """Same sheet enumeration and row-to-text strategy as Dify 1.10.1 ExcelExtractor."""
    import pandas as pd
    workbook=load_workbook(io.BytesIO(blob),data_only=True)
    output=[]
    for sheet in workbook:
        rows=sheet.values;columns=next(rows,None)
        if columns is None:continue
        frame=pd.DataFrame(rows,columns=columns).dropna(how='all')
        for _,row in frame.iterrows():
            output.append(';'.join('"%s":"%s"'%(key,value) for key,value in row.items() if pd.notna(value)))
    return '\n'.join(output)


def priced_example(price=None,unit='只',quantity=1):
    products=[row for row in real_candidates()[0] if row['material_code']=='335012800535']
    prices=price or json.loads((FIXTURES/'live_qf_price.json').read_text())
    item={'material_code':'335012800535','material_name':'小型断路器','model_spec':'TGB1N-125H 4P C100',
          'unit':unit,'quantity':quantity,'requirements_json':'{}','key_technical_parameters':'',
          'source_request_id':'REQ-001','quantity_basis':'PROJECT'}
    ns=load_node('3400000000005')
    class FixedDatetime(datetime):
        @classmethod
        def now(cls,tz=None):return cls(2026,10,28,12,tzinfo=timezone.utc).astimezone(tz)
    ns['datetime']=FixedDatetime
    result=ns['main'](json.dumps(item,ensure_ascii=False),item['material_code'],products,prices,products,prices,'HIT','HIT')
    return json.loads(result['quote_item_json'])
