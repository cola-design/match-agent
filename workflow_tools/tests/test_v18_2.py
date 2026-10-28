import ast
from copy import deepcopy
import hashlib
import io
import json
import unittest
from openpyxl import load_workbook
from case_helpers import DOC,NODES,ROOT,load_node,case_requirements,real_candidates,selected_case,fill_output,dify_extract,priced_example


class V182Test(unittest.TestCase):
    def test_graph_and_dynamic_template_selectors(self):
        self.assertEqual(len(NODES),108)
        for edge in DOC['workflow']['graph']['edges']:
            self.assertIn(edge['source'],NODES);self.assertIn(edge['target'],NODES)
            self.assertEqual(edge['data']['sourceType'],NODES[edge['source']]['data']['type'])
            self.assertEqual(edge['data']['targetType'],NODES[edge['target']]['data']['type'])
        for node in NODES.values():
            if node['data']['type']=='code':compile(node['data']['code'],node['id'],'exec')
            if node['data']['type']=='tool':
                selector=node['data']['tool_configurations']['template_base64']
                self.assertEqual(selector['type'],'variable')
                self.assertEqual(node['data']['tool_node_version'],'2')
                self.assertIn(selector['value'][1],NODES[selector['value'][0]]['data']['outputs'])
        for variable in DOC['workflow']['environment_variables']:
            if variable['name'].startswith('MONGODB_'):self.assertEqual(variable['value'],'')

    def test_product_and_price_contracts_reject_stringified_or_invalid_fields(self):
        product=real_candidates()[0][0]
        normalize=load_node('2000000000022')['main']
        self.assertEqual(normalize(json.dumps([product]),200)['result'][0]['material_code'],'335012800535')
        with self.assertRaises(ValueError):normalize(json.dumps([{**product,'category':json.dumps(product['category'])}]),200)
        price=load_node('3400000000004')['main']
        with self.assertRaises(ValueError):price('[{"material_code":"001","unit_price_ex_tax":true,"tax_rate":0.13}]',200)
        for status in (401,503):
            with self.assertRaises(ValueError):normalize('[]',status)

    def test_source_units_and_protocol_evidence_survive_auxiliary_merge(self):
        rows,issues=case_requirements();self.assertEqual(issues,[])
        self.assertEqual([row['unit'] for row in rows],['只','只','只','只','个','个','个','件','件','件','个','套'])
        for row in rows:
            self.assertEqual(row['mapping_status'],'MATCHED')
            self.assertTrue(row['source_evidence']['protocol'])
            self.assertEqual(row['unit_basis'],'VERIFIED_DOCUMENT')
        self.assertEqual(rows[-1]['requirements']['door_count'],2)
        self.assertEqual(rows[-1]['requirements']['door_type'],'双门')

    def test_unit_conflict_and_missing_unit_do_not_get_silent_defaults(self):
        root=load_node('1800000000007')
        ev={'text':'开关1只','source_verified':True,'document_hash':'x','chunk_id':'c'}
        for drawing in ([{**ev,'text':'开关1套'}],[]):
            row={'request_id':'r','unit':'台','requirements':{},'source_evidence':{'drawing':drawing,'protocol':[ev] if drawing else []}}
            issues=[];root['preserve_source_units_and_structure']({'component_requirements':[row]}, {},{},issues)
            self.assertTrue(any(issue['blocks_retrieval'] for issue in issues))
        self.assertEqual(root['explicit_source_units']([{**ev,'text':'开关一只'}]),{'只'})
        self.assertEqual(root['explicit_source_units']([{**ev,'text':'单位：件'}]),{'件'})
        self.assertEqual(root['explicit_source_units']([{**ev,'text':'304 M8×80，四个'}]),{'个'})

    def test_embedded_clarification_validator_matches_root_validator(self):
        source=NODES['3600000000008']['data']['code'];tree=ast.parse(source)
        assignment=next(n for n in tree.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='_REQUIREMENT_VALIDATOR_CODE' for t in n.targets))
        self.assertEqual(ast.literal_eval(assignment.value),NODES['1800000000007']['data']['code'])

    def test_actual_search_variants_require_explicit_accessory_choice(self):
        rows,_=case_requirements();req=rows[0];ns=load_node('2000000000024')
        out=ns['main']('',json.dumps(req,ensure_ascii=False),real_candidates()[0])
        result=json.loads(out['item_json'])
        self.assertEqual(result['match_status'],'NEED_REVIEW')
        self.assertFalse(result['recommended'])
        self.assertEqual({r['material_code'] for r in result['alternatives']},{'335012800535','335012800536','335012800537'})
        req['requirements']['accessories_explicit_none']=True
        checked=ns['main']('',json.dumps(req,ensure_ascii=False),real_candidates()[0])
        result=json.loads(checked['item_json'])
        self.assertEqual(result['recommended']['material_code'],'335012800535')
        self.assertEqual(result['match_status'],'MATCHED')

    def test_search_model_is_generic_and_keeps_bare_size_out_of_model_filter(self):
        query=load_node('2000000000021')['main']
        for model,wanted in [('CJX2-3210','CJX2-3210'),('TGM1-250L','TGM1-250L'),('270',''),('550×700×150mm','')]:
            out=query({'request_id':'r','shown_model':model,'component_role':'器件','requirements':{}})
            self.assertEqual(out['search_model'],wanted)
        out=query({'request_id':'r','shown_model':'TGB1N-125H','requirements':{'rated_current_a':100.0,'trip_curve':'C','pole_count':'4P'}})
        self.assertIn('C100',out['retrieval_query']);self.assertNotIn('C100.0',out['retrieval_query'])

    def test_non_breaker_family_respects_current_poles_and_control_voltage(self):
        selector=load_node('2000000000024')['main']
        product={'material_code':'00123456','description':'CJX2-3210 3P 32A AC220V',
                 'category':{'middle':'接触器','leaf':'CJX2'},'standard_specs':{},
                 'source_attributes':[{'name':'额定电流','value':'32A'},{'name':'极数','value':'3P'},{'name':'线圈电压','value':'AC220V'}],
                 'specification_text':'额定电流=32A；极数=3P；线圈电压=AC220V','source_status':'有效','data_status':'ACTIVE','quality_flags':[]}
        req={'request_id':'r','component_role':'交流接触器','shown_model':'CJX2-3210','mapping_status':'MATCHED',
             'requirements':{'rated_current_a':32,'pole_count':'3P','control_voltage':'AC220V'},'quantity':1,'unit':'只'}
        result=json.loads(selector('',json.dumps(req),[product])['item_json'])
        self.assertEqual(result['match_status'],'MATCHED',result.get('warnings'))
        for field,value in [('rated_current_a',63),('pole_count','4P'),('control_voltage','DC220V')]:
            changed=deepcopy(req);changed['requirements'][field]=value
            result=json.loads(selector('',json.dumps(changed),[product])['item_json'])
            self.assertEqual(result['match_status'],'NOT_RETRIEVED',(field,result))

    def test_two_scheme_per_cabinet_rows_remain_separate(self):
        rows,_=case_requirements();source=rows[1]
        candidates=[p for p in real_candidates()[1] if 'TGH1N-125' in p['description']]
        selector=load_node('2000000000024');requirements=[];outputs=[]
        for i in (1,2):
            row=deepcopy(source);row.update(request_id='SCOPE-'+str(i),scheme_group_id='D-'+str(i)*16,quantity_basis='PER_CABINET',quantity_per_cabinet=1,project_quantity=None,quantity=None,project_cabinet_quantity=None)
            requirements.append(row);outputs.append(selector['main']('',json.dumps(row),candidates)['item_json'])
        aggregate=load_node('2000000000025');out=aggregate['main'](outputs,aggregate['wire_pack'](requirements))
        context=aggregate['wire_loads'](out['selection_json'])
        self.assertEqual(len(context['items']),2)
        self.assertTrue(all(row['quantity_basis']=='PER_CABINET' and row['project_quantity'] is None for row in context['items']))
        self.assertNotEqual(context['items'][0]['position_display'],context['items'][1]['position_display'])

    def test_case_retains_all_12_rows_and_never_fills_missing_auxiliary_codes(self):
        out,rows=selected_case();ns=load_node('2000000000025');context=ns['wire_loads'](out['selection_json'])
        self.assertEqual(len(context['items']),12)
        self.assertFalse(context['pricing_allowed'])
        for item in context['items'][4:]:self.assertFalse(item.get('recommended'))
        self.assertNotIn('C-71d0abbee75a7249',out['confirm_markdown'])
        self.assertIn('需求：PG29',out['confirm_markdown'])
        self.assertIn('1 只',out['confirm_markdown'])

    def test_dynamic_workbook_preserves_complete_context_and_print_geometry(self):
        selected,_=selected_case();writer=load_node('1800000000103')
        out=writer['main'](selected['confirm_markdown'],selected['selection_json'])
        blob=fill_output(out);workbook=load_workbook(io.BytesIO(blob))
        main=workbook['BOM确认']
        self.assertIn('$J$16',str(main.print_area))
        self.assertEqual(main.print_title_rows,'$1:$4')
        self.assertEqual(workbook['系统数据'].sheet_state,'veryHidden')
        self.assertTrue(main.row_dimensions[130].hidden)
        self.assertEqual(main.print_options.horizontalCentered,True)
        self.assertEqual(main.page_setup.fitToWidth,1)
        self.assertEqual(main.page_setup.fitToHeight,0)
        self.assertTrue(all(34<=main.row_dimensions[row].height<=409 for row in range(5,17)))
        self.assertFalse(main['D17'].has_style)
        extracted=dify_extract(blob)
        restored=load_node('3300000000022')['stored_selection'](extracted)
        self.assertEqual(restored,writer['wire_loads'](selected['selection_json']))
        guard=load_node('3300000000022')
        headers=guard['HEADERS']
        rows=[{key:main.cell(number,index+1).value or '' for index,key in enumerate(headers)} for number in range(5,17)]
        opinions=[{'序号':str(index),'normalized_action':'PENDING','action_confidence':1,'action_reason':'未填写'} for index in range(1,13)]
        guarded=guard['main'](json.dumps({'form_type':'BOM_CONFIRMATION','rows':rows},ensure_ascii=False),json.dumps({'form_type':'BOM_CONFIRMATION','rows':opinions},ensure_ascii=False),extracted)
        self.assertEqual(guarded['guard_status'],'VALID',guarded['guard_message'])
        rows[0]['关键技术参数']='额定电流：125A'
        guarded=guard['main'](json.dumps({'form_type':'BOM_CONFIRMATION','rows':rows},ensure_ascii=False),json.dumps({'form_type':'BOM_CONFIRMATION','rows':opinions},ensure_ascii=False),extracted)
        self.assertEqual(guarded['guard_status'],'INVALID')

    def test_price_contract_preserves_real_source_and_marks_estimate(self):
        row=priced_example()
        self.assertEqual(row['item_status'],'PRICED')
        self.assertEqual(row['unit_price_ex_tax'],60)
        self.assertEqual(row['tax_rate'],0.13)
        self.assertTrue(row['reference_only'])
        self.assertIn('只',row['unit_conversion_note'])

    def test_quote_precision_config_and_price_change_identifiers(self):
        ns=load_node('3400000000006');row=priced_example()
        first=ns['main']([row],markup_rate='0',discount_rate='0')
        self.assertEqual(first['material_subtotal'],60)
        self.assertEqual(first['tax_included_quote'],67.8)
        self.assertEqual(first['reference_price_count'],1)
        self.assertIn('参考',first['quote_title'])
        second=ns['main']([{**row,'unit_price_ex_tax':61}],markup_rate='0',discount_rate='0')
        self.assertNotEqual(first['quote_no'],second['quote_no'])
        third=ns['main']([{**row,'quantity':10000,'unit_price_ex_tax':0.0149,'tax_rate':0}],markup_rate='0',discount_rate='0')
        self.assertEqual(third['material_subtotal'],149)
        self.assertIn('BOM数量单位',json.loads(first['quote_payload_json'])['cells']['02 报价明细!L4'])
        workbook=load_workbook(io.BytesIO(fill_output(first,'quote_payload_json')))
        self.assertIn('ROUND',workbook['02 报价明细']['H4'].value)
        self.assertIn('$L$4',str(workbook['02 报价明细'].print_area))

    def test_partial_results_and_invalid_configuration_block(self):
        ns=load_node('3400000000006');row=priced_example()
        out=ns['main']([row,{'quantity':1,'item_status':'PRICE_NOT_RETRIEVED','issue':'missing'}])
        self.assertEqual(out['pricing_status'],'BLOCKED');self.assertEqual(out['material_subtotal'],0)
        with self.assertRaises(ValueError):ns['main']([row],markup_rate='NaN')
        with self.assertRaises(ValueError):ns['main']([row],discount_rate='2')

    def test_dynamic_layout_handles_one_and_120_compact_rows(self):
        writer=load_node('1800000000103')
        for count in (1,120):
            context={'items':[{'request_id':str(i),'match_status':'MATCHED','warnings':[],'alternatives':[]} for i in range(1,count+1)],'pricing_allowed':False}
            header='## BOM确认\n\n| 序号 | 回路/位号 | 元器件名称 | 关键技术参数 | 推荐型号 | 物料编码 | 数量 | 系统状态 | 工程师确认 | 修改意见 |\n|---:|---|---|---|---|---|---|---|---|---|\n'
            table='\n'.join('| %d | X%d | 器件 | 参数 | 待选 |  | 1 件 | 待确认 |  |  |'%(i,i) for i in range(1,count+1))
            out=writer['main'](header+table,json.dumps(context))
            book=load_workbook(io.BytesIO(fill_output(out)))
            self.assertIn('$J$'+str(count+4),str(book['BOM确认'].print_area))
            self.assertEqual(book['选型说明'].sheet_state,'hidden')
            self.assertTrue(book['BOM确认'].row_dimensions[130].hidden)

    def test_unshowable_long_cell_stops_instead_of_silent_clipping(self):
        writer=load_node('1800000000103')
        context={'items':[{'request_id':'r','match_status':'MATCHED','warnings':[],'alternatives':[]}],'pricing_allowed':False}
        markdown='## BOM确认\n\n| 1 | X1 | 器件 | '+'长参数'*700+' | 待选 |  | 1 件 | 待确认 |  |  |'
        with self.assertRaisesRegex(ValueError,'显示上限'):
            writer['main'](markdown,json.dumps(context))


if __name__=='__main__':unittest.main()
