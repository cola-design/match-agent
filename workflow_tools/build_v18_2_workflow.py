"""Build V18.2 from the V18.1 workflow. No live credentials or database access."""
import argparse
import ast
import base64
from copy import copy,deepcopy
import io
import json
from pathlib import Path
import re
import uuid
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Font,PatternFill,Border,Side

ROOT=Path(__file__).resolve().parents[1]
RUNTIME=ROOT/'workflow_tools/runtime'
BLUE='17365D';LIGHT='EAF2F8';YELLOW='FFF2CC';LINE='C8D6E5'
LAYOUTS={
 'bom':{'BOM确认':dict(start=5,columns=10,widths=[6,13,21,44,27,19,12,17,17,25],header_end=4),
        '选型说明':dict(start=4,columns=5,widths=[13,25,19,72,48],header_end=3,hide_if_empty=True)},
 'clarification':{'资料问题确认':dict(start=5,columns=6,widths=[7,22,18,66,44,45],header_end=4)},
 'drawing':{'识别概况':dict(start=2,columns=2,widths=[25,100],header_end=1,paper=9),
            '待处理问题':dict(start=2,columns=4,widths=[7,15,75,55],header_end=1),
            '已识别方案组':dict(start=2,columns=10,widths=[7,16,26,10,10,14,15,13,14,48],header_end=1)},
 'quote':{'01 报价汇总':dict(start=6,minimum=21,columns=4,widths=[28,28,18,18],header_end=5,paper=9,orientation='portrait',min_height=24),
          '02 报价明细':dict(start=4,columns=12,widths=[6,18,22,29,9,8,13,14,10,15,20,45],header_end=3),
          '03 缺价与说明':dict(start=4,columns=6,widths=[7,19,24,22,58,42],header_end=3)},
}


def replace_once(source,old,new):
    if source.count(old)!=1:raise ValueError('Expected exactly one source fragment: '+old[:100])
    return source.replace(old,new,1)


def styled_template(raw,kind):
    workbook=load_workbook(io.BytesIO(raw))
    if kind=='bom':
        details=workbook.create_sheet('选型说明')
        details.merge_cells('A1:E1');details['A1']='选型核对说明与候选差异'
        details.append([]);details.append(['位号','资料要求型号','匹配状态','核对说明','候选型号及编码'])
        for row in details.iter_rows(min_row=4,max_row=123,max_col=5):
            for cell in row:cell.fill=PatternFill('solid',fgColor=LIGHT)
        workbook['BOM确认']['A2']='蓝色区域为来源要求与候选结果，黄色区域填写工程师意见；详细原因和候选差异见“选型说明”。'
        workbook['BOM确认']['A3']='“需求：”仅表示资料指定型号，不代表已找到企业编码。请只填写 I、J 列；有候选差异时写明最终型号和编码。'
    if kind=='quote':
        summary=workbook['01 报价汇总'];summary['A2']='成套柜材料报价'
        summary['A3']='价格依据、版本及有效期见明细；含估算价格时输出材料参考报价。'
        summary['A18']='范围：仅包含本次已确认BOM中的材料费用。'
        summary['A19']='未包含人工、加工、检验、运输等费用；不等于整柜项目总价。'
        summary['A20']='估算价格不等于已核实采购价；真实价格按版本和有效期维护。'
        summary['A21']='加成及折扣按本次配置列示，正式使用须按企业审批规则执行。'
        summary['A10']='调整后未税金额';summary['A11']='调整后含税金额';summary['B13']='本次配置'
        for row in (5,6,7,8,9,10,11,13,14,15,16):summary.merge_cells(start_row=row,start_column=2,end_row=row,end_column=4)
        workbook['03 缺价与说明']['A2']='价格依据及报价范围说明'
        workbook['02 报价明细']['A2']='材料计价明细（全单加成／折扣调整前）'
        detail=workbook['02 报价明细']
        for row in range(4,124):
            detail.cell(row,7).number_format='0.00####'
            detail.cell(row,8).value=f'=IF(OR(E{row}="",G{row}=""),"",ROUND(E{row}*G{row},2))'
            detail.cell(row,10).value=f'=IF(OR(H{row}="",I{row}=""),"",ROUND(H{row}*(1+I{row}),2))'
    border=Border(*(Side(style='thin',color=LINE) for _ in range(4)))
    for sheet in workbook:
        if sheet.title=='系统数据':sheet.sheet_state='veryHidden';continue
        config=LAYOUTS[kind][sheet.title]
        for number,width in enumerate(config['widths'],1):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(number)].width=width
        for row in sheet:
            for cell in row:
                if cell.__class__.__name__=='MergedCell':continue
                font=copy(cell.font);font.name='Microsoft YaHei';font.sz=10.5
                if cell.row<config['start'] and font.bold:font.sz=11
                cell.font=font
                alignment=copy(cell.alignment);alignment.wrap_text=True;alignment.vertical='top'
                cell.alignment=alignment
                if cell.row>=config['start'] and cell.row<130:cell.border=border
        if sheet.title=='选型说明':
            sheet['A1'].fill=PatternFill('solid',fgColor=BLUE);sheet['A1'].font=Font(name='Microsoft YaHei',size=17,bold=True,color='FFFFFF')
            sheet.row_dimensions[1].height=34
            for cell in sheet[3]:cell.fill=PatternFill('solid',fgColor=BLUE);cell.font=Font(name='Microsoft YaHei',size=11,bold=True,color='FFFFFF')
        if kind in ('bom','clarification') and sheet.title!='选型说明':
            for number,height in ((1,34),(2,32),(3,38),(4,30)):sheet.row_dimensions[number].height=height
            sheet['A1'].font=Font(name='Microsoft YaHei',size=18,bold=True,color='FFFFFF')
        if sheet.title=='01 报价汇总':
            sheet['A2'].font=Font(name='Microsoft YaHei',size=18,bold=True,color='FFFFFF')
            sheet.row_dimensions[2].height=38;sheet.row_dimensions[3].height=30
        sheet.freeze_panes='A'+str(config['start'])
        sheet.sheet_view.showGridLines=False
        sheet.print_options.horizontalCentered=True
        sheet.oddFooter.center.text='第 &P 页 / 共 &N 页'
        sheet.oddFooter.center.size=9
        if sheet.title!='01 报价汇总':sheet.auto_filter.ref='A'+str(config['start']-1)+':'+get_column_letter(config['columns'])+'123'
    output=io.BytesIO();workbook.save(output)
    return output.getvalue()


def patch_requirement_source(source):
    fragment="peer.setdefault('protocol_requirements_evidence', []).append(copy.deepcopy(req))"
    source=replace_once(source,fragment,fragment+"\n                        restore_peer_protocol_evidence(peer, req)")
    fragment="payload = preserve_cabinet_context(payload, drawing, protocol)\n        out = publish_checked_output(out, payload, issues)"
    source=replace_once(source,fragment,"payload = preserve_cabinet_context(payload, drawing, protocol)\n        payload = preserve_source_units_and_structure(payload, drawing, protocol, issues)\n        out = publish_checked_output(out, payload, issues)")
    return source+'\n'+(RUNTIME/'source_fixes.py').read_text()


def patch_embedded_validator(source,validator):
    tree=ast.parse(source)
    node=next(n for n in tree.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='_REQUIREMENT_VALIDATOR_CODE' for t in n.targets))
    old=ast.get_source_segment(source,node)
    return replace_once(source,old,'_REQUIREMENT_VALIDATOR_CODE = '+repr(validator))


NORMALIZER='''import json, math
def main(body: str, status_code: int) -> dict:
    if int(status_code)!=200:
        raise ValueError('API_HTTP_%s：请检查API地址、密钥、网络、数据库连接和文档结构；本次停止。'%status_code)
    try:rows=json.loads(body)
    except (TypeError,ValueError):raise ValueError('API_RESPONSE_INVALID_JSON')
    if not isinstance(rows,list) or any(not isinstance(row,dict) for row in rows):raise ValueError('API_RESPONSE_NOT_DOCUMENT_ARRAY')
    for row in rows:
        if not isinstance(row.get('material_code'),str) or not row['material_code'].strip():raise ValueError('API_MATERIAL_CODE_MUST_BE_STRING')
        CONTRACT
    return {'result':rows}
'''
PRODUCT_CONTRACT="""for field,expected in (('category',dict),('standard_specs',dict),('source_attributes',list),('quality_flags',list)):
            if not isinstance(row.get(field),expected):raise ValueError('API_PRODUCT_FIELD_TYPE_INVALID:'+field)"""
PRICE_CONTRACT="""for field in ('unit_price_ex_tax','tax_rate'):
            value=row.get(field)
            if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):raise ValueError('API_PRICE_FIELD_TYPE_INVALID:'+field)
        for field in ('pricing_unit','price_version','valid_from','valid_to','status'):
            if not isinstance(row.get(field),str) or not row[field].strip():raise ValueError('API_PRICE_FIELD_MISSING:'+field)"""


def build(source_path=None):
    path=source_path or ROOT/'formal_selection_quote_v18_1_http_api.yml'
    document=yaml.safe_load(path.read_text())
    nodes={n['id']:n for n in document['workflow']['graph']['nodes']}
    document['app']['name']='成套柜智能选型与报价 V18.2（API与动态排版）'
    document['app']['description']='保留V18.1模型、人工确认和API链路；修正来源单位、协议证据、候选差异、估算价格标识，按实际行数与内容动态生成Excel模板。'
    for variable in document['workflow'].get('environment_variables',[]):variable['value']=''
    for name,value,label in [('MATERIAL_MARKUP_RATE','0.15','材料报价加成率；0.15表示15%，请按业务核准配置'),('QUOTE_DISCOUNT_RATE','0.02','报价折扣率；0.02表示减2%，请按业务核准配置')]:
        document['workflow']['environment_variables'].append({'id':str(uuid.uuid5(uuid.NAMESPACE_URL,'match-agent/v18.2/'+name)),'name':name,'value_type':'string','value':value,'description':label,'selector':['env',name]})
    validator=patch_requirement_source(nodes['1800000000007']['data']['code'])
    nodes['1800000000007']['data']['code']=validator
    nodes['3600000000008']['data']['code']=patch_embedded_validator(nodes['3600000000008']['data']['code'],validator)
    for nid in ('2000000000022','3400000000003','3400000000009'):
        nodes[nid]['data']['code']=NORMALIZER.replace('CONTRACT',PRODUCT_CONTRACT)
    for nid in ('3400000000004','3400000000010'):
        nodes[nid]['data']['code']=NORMALIZER.replace('CONTRACT',PRICE_CONTRACT)
    for node in nodes.values():
        if node['data']['type']=='http-request':node['data']['retry_config']={'max_retries':2,'retry_enabled':True,'retry_interval':1000}
        if node['data']['type']=='code':
            node['data']['code']=node['data']['code'].replace('请调整知识库分段/检索词或补充产品数据','请核对API检索词、候选过滤与产品发布状态，或补充产品数据')
            node['data']['code']=node['data']['code'].replace('PRICE知识库绑定/检索','价格API响应及有效期')
        if node['data']['type']=='llm':
            prompts=node['data'].get('prompt_template',[])
            for prompt in prompts:
                if prompt.get('role')=='system' and 'cabinet_material' in prompt.get('text',''):
                    prompt['text']+='\n柜体单门/双门等结构要求仅在资料明确时提取：方案组保存cabinet_door_count、cabinet_door_type，柜体条目shown_parameters保存door_count、door_type；对应方案组或条目evidence必须包含原文。计量单位从对应条目引文提取，不默认统一为台/个；没有明确单位留空。'
    query=nodes['2000000000021']['data']
    query['code']+='''\n_v182_query_base=main
import unicodedata
from decimal import Decimal
_v182_first_value=first_value
def first_value(data,*keys):
    value=_v182_first_value(data,*keys)
    return format(Decimal(value).normalize(),'f') if re.fullmatch(r'[0-9]+(?:\\.[0-9]+)?',value) else value
def main(requirement):
    out=_v182_query_base(requirement)
    item=wire_loads(out['requirement_json'])
    model=unicodedata.normalize('NFKC',clean(item.get('shown_model')))
    token=re.split(r'[\\s;；,，]+',model)[0]
    out['search_model']=token if re.fullmatch(r'(?=.*[A-Za-z])(?=.*\\d)[A-Za-z][A-Za-z0-9._/-]{2,}',token) else ''
    out['search_category']=''
    return out
'''
    query['outputs'].update(search_model={'type':'string','children':None},search_category={'type':'string','children':None})
    nodes['1830000000001']['data']['params']='q:{{#2000000000021.retrieval_query#}}\nmodel:{{#2000000000021.search_model#}}\nlimit:30'
    selector=nodes['2000000000024']['data']
    selector['code']+=(RUNTIME/'selection_fixes.py').read_text()
    aggregate=nodes['2000000000025']['data']
    aggregate['code']=patch_display(aggregate['code'])
    price=nodes['3400000000005']['data']
    price['code']=price['code'].replace('date.today()', 'datetime.now(timezone(timedelta(hours=8))).date()')
    price['code']+='\nfrom datetime import timezone, timedelta\n'
    price['code']+=(RUNTIME/'price_fixes.py').read_text()
    quote=nodes['3400000000006']['data']
    quote['code']=patch_quote(quote['code'])
    quote['variables'] += [{'variable':name,'value_type':'string','value_selector':['env',env]} for name,env in [('markup_rate','MATERIAL_MARKUP_RATE'),('discount_rate','QUOTE_DISCOUNT_RATE')]]
    quote['outputs'].update(quote_title={'type':'string','children':None},reference_price_count={'type':'number','children':None})
    for tid,pid,kind in [('1900000000014','1800000000103','bom'),('1800000000014','1800000000102','clarification'),('1800000000013','1800000000101','drawing'),('3300000000017','3400000000006','quote')]:
        tool=nodes[tid]['data'];config=tool['tool_configurations']['template_base64']
        raw=base64.b64decode(config['value'] if isinstance(config,dict) else config)
        raw=styled_template(raw,kind)
        template=ROOT/'workflow_templates'/('v18_2_'+kind+'.xlsx');template.write_bytes(raw)
        code=nodes[pid]['data']
        if kind=='bom':code['code']+=(RUNTIME/'bom_payload.py').read_text()
        code['code']+=layout_wrapper(base64.b64encode(raw).decode(),kind,'quote_payload_json' if kind=='quote' else 'payload_json')
        code['outputs']['template_base64']={'type':'string','children':None}
        tool['tool_configurations']['template_base64']={'type':'variable','value':[pid,'template_base64']}
        tool['tool_configurations']['output_filename']=({'type':'variable','value':[pid,'output_filename']} if kind=='quote' else {'type':'constant','value':{'bom':'BOM人工确认单_V18.2','clarification':'资料问题确认单_V18.2','drawing':'图纸识别问题确认单_V18.2'}[kind]})
        tool['tool_node_version']='2'
        tool['desc']='使用上一代码节点按真实行数和文本长度生成的动态模板；保留蓝黄配色、人工确认列和无损回传上下文。'
    answer=nodes['3300000000018']['data']
    nodes['3300000000017']['data']['title']='生成材料报价Excel（按来源标识）'
    answer['title']='输出材料报价Excel'
    answer['answer']='''## {{#3400000000006.quote_title#}}

{{#3400000000006.pricing_message#}}

- 报价编号：{{#3400000000006.quote_no#}}
- 已计价物料：{{#3400000000006.priced_count#}} 项
- 材料未税小计：{{#3400000000006.material_subtotal#}} 元
- 按配置调整后未税金额：{{#3400000000006.pre_tax_quote#}} 元
- 按配置调整后含税金额：{{#3400000000006.tax_included_quote#}} 元

{{#3300000000017.files#}}

价格依据、版本、有效期和单位换算见明细。含估算价格时仅作材料参考；未计人工、加工、检验与运输，不等于整柜项目总价。'''
    output=ROOT/'formal_selection_quote_v18_2_http_api.yml'
    for node in nodes.values():
        if node['data']['type']=='code':compile(node['data']['code'],node['id'],'exec')
    output.write_text(yaml.safe_dump(document,allow_unicode=True,sort_keys=False,width=1000))
    assert yaml.safe_load(output.read_text())==document
    return output


def patch_display(source):
    source=source.replace("'workflow_version'] = 'V17.3'","'workflow_version'] = 'V18.2'").replace('版本：V17.3','版本：V18.2')
    source=replace_once(source,"position = '/'.join((x for x in ([clean(item.get('scheme_group_id'))] if many_groups else []) + [clean(item.get('symbol')), clean(item.get('circuit_id'))] if x))", "position = v182_position(item, normalized, many_groups)")
    source=replace_once(source,"visible_status += '：' + review_reasons","visible_status += '（见选型说明）'")
    source=replace_once(source,"model_name = clean(rec.get('material_name') or rec.get('product_model') or rec.get('model') or rec.get('series')) or source_model", "model_name = clean(rec.get('material_name') or rec.get('product_model') or rec.get('model') or rec.get('series')) or ('需求：' + source_model if source_model else '未匹配')")
    source=replace_once(source,"model_name = '待工程师选定：' + '；'.join(choices)","model_name = '待选定（见选型说明）'")
    source=replace_once(source,"visible_status = '待选定型号：' + review_reasons","visible_status = '待选定（见选型说明）'")
    source=replace_once(source,"requirement = clean(item.get('requirement_summary'))","requirement = v182_requirement_display(item.get('requirement_summary'))\n        item['requirement_summary'] = requirement")
    return source+(RUNTIME/'display_fixes.py').read_text()


def patch_quote(source):
    source=replace_once(source,"price = money(item.get('unit_price_ex_tax'))", "price = Decimal(str(item.get('unit_price_ex_tax')))")
    source=replace_once(source,'subtotal += qty * price','subtotal += money(qty * price)')
    source=source.replace("'01 报价汇总!B14': 0.15","'01 报价汇总!B14': float(_V182_MARKUP)")
    source=source.replace("'01 报价汇总!B15': 0.02","'01 报价汇总!B15': float(_V182_DISCOUNT)")
    source=source.replace("subtotal * Decimal('1.15') * Decimal('0.98')","subtotal * (Decimal('1')+_V182_MARKUP) * (Decimal('1')-_V182_DISCOUNT)")
    source=source.replace("fingerprint = '|'.join((clean(x.get('material_code')) + ':' + clean(x.get('quantity')) for x in items))", "fingerprint = json.dumps({'items': [{k:x.get(k) for k in ('material_code','quantity','unit_price_ex_tax','tax_rate','price_version','price_validity','price_basis')} for x in items], 'markup':str(_V182_MARKUP),'discount':str(_V182_DISCOUNT)},ensure_ascii=False,sort_keys=True)")
    source=source.replace('按报价规则15%%加成、2%%折扣计价','按本次配置的加成与折扣计价（比例见汇总表）')
    return source+(RUNTIME/'quote_fixes.py').read_text()


def layout_wrapper(template,kind,key):
    runtime=(RUNTIME/'excel_layout.py').read_text()
    argument_names={'bom':('markdown, selection_json','markdown, selection_json'), 'clarification':('markdown, selection_request_json','markdown, selection_request_json'),
                    'drawing':('markdown','markdown'),'quote':("results, confirmed_bom_json='', markup_rate='0.15', discount_rate='0.02'",'results, confirmed_bom_json, markup_rate, discount_rate')}
    signature,args=argument_names[kind]
    code='\n'+runtime+'\n_LAYOUT_TEMPLATE='+repr(template)+'\n_LAYOUT_SPEC='+repr(LAYOUTS[kind])+'\n_v182_layout_base=main\n'
    code+='def main('+signature+'):\n    out=_v182_layout_base('+args+')\n'
    code+='    out["template_base64"]=layout_template(_LAYOUT_TEMPLATE,out['+repr(key)+'],_LAYOUT_SPEC)\n'
    code+='    return out\n'
    return code


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--source',type=Path)
    print(build(parser.parse_args().source))
